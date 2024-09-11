'''modules'''
import os
import openpyxl
from docx import Document

def replace_placeholders(doc, row_values):
    '''document placeholder replacement based on column indices'''

    for paragraph in doc.paragraphs:
        for i, value in enumerate(row_values):
            placeholder = f'<<<{i}>>>'

            if placeholder in paragraph.text:
                paragraph.text = paragraph.text.replace(placeholder, value)


def process_excel_file(excel_path, template_path, output_dir):
    '''Excel file process and documents generation'''
    sheet = openpyxl.load_workbook(excel_path).active

    for i, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=1):
        doc = Document(template_path)
        row_values = [str(value) if value is not None else '' for value in row]

        output_path = os.path.join(
            output_dir,
            f'{'_'.join([part if part else 'unknown' for part in row_values[:3]])}.docx'
            )

        replace_placeholders(doc, row_values)
        doc.save(output_path)

        print(f'File {i} from {os.path.basename(excel_path)} saved in: {output_path}')


def get_excel_files(directory):
    '''folder xlsx files search'''

    return [
        os.path.join(directory, file)
        for file in os.listdir(directory)
        if file.endswith('.xlsx')
        ]


if __name__ == '__main__':
    EXCEL_FOLDER = 'data'
    TEMPLATE_PATH = 'template.docx'
    OUTPUT_DIR = 'results'

    if not os.path.exists(EXCEL_FOLDER):
        os.makedirs(EXCEL_FOLDER)

    if not os.path.exists(OUTPUT_DIR):
        os.makedirs(OUTPUT_DIR)

    if not os.path.exists(TEMPLATE_PATH):
        raise FileNotFoundError(f'No file {TEMPLATE_PATH}')

    excel_files = get_excel_files(EXCEL_FOLDER)

    if not excel_files:
        raise FileNotFoundError(f'No xlsx files in folder: {EXCEL_FOLDER}')

    for excel_file in excel_files:
        process_excel_file(excel_file, TEMPLATE_PATH, OUTPUT_DIR)
